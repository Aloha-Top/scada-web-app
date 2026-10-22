"""
app.py - ระบบ Web Server, REST API และ Background Scheduler สำหรับระบบ Web Portal, ระบบจัดซื้อจัดจ้าง และแผนที่/แดชบอร์ด SCADA (%AVA)
ออกแบบมาเพื่อทำงานร่วมกับไฟล์ index.html, procurement.html, create_detailed_map.py และรองรับการ Deploy บนแพลตฟอร์ม Render

จุดเด่นและการปรับปรุงหลัก:
1. Web Server & Render Deployment:
   - ใช้ Flask เป็น Backend Framework หลัก (พร้อม Fallback HTTP Server กรณีรันเครื่อง Local ที่ยังไม่ได้ติดตั้ง Flask)
   - รันด้วย host="0.0.0.0" และอ่านค่าพอร์ตจาก Environment Variable `os.environ.get('PORT', 10000)`
2. Background Task & In-Memory Cache (ป้องกัน Google Sheets Rate Limit):
   - ใช้ `APScheduler` (`BackgroundScheduler`) ตั้งเวลาดึงข้อมูลจาก Google Sheets และประมวลผลอัตโนมัติทุกๆ 5 นาที (ปรับค่าได้ผ่าน `SCADA_REFRESH_MINUTES`)
   - เก็บผลลัพธ์ทั้งหมดไว้ในตัวแปรหน่วยความจำ (`GLOBAL_SCADA_CACHE`) เพื่อให้ทุก API เรียกใช้ข้อมูลได้ทันทีโดยไม่ต้องยิงตรงไปที่ Google Sheets ทุกครั้งที่มีคนเปิดหน้าเว็บ
3. Web Portal & API Endpoints:
   - `/` (และ `/index.html`, `/portal`): เสิร์ฟหน้า Web Portal หลักของหน่วยงาน (index.html)
   - `/procurement.html` (และ `/procurement`): เสิร์ฟหน้าระบบบริหารและติดตามงานจัดซื้อจัดจ้าง (procurement.html)
   - `/map.html` และ `/dashboard.html` (รวมถึง `/map`, `/dashboard`, `/ai-report`): เสิร์ฟหน้าแผนที่และแดชบอร์ด SCADA (%AVA)
   - `/api/map_data` (และ `/api/map-data`): คืนค่า JSON สำหรับอัปเดตข้อมูลแผนที่และแดชบอร์ดแบบ Seamless Update
"""

import io
import json
import os
import pathlib
import sys
import threading
import time

try:
    from dotenv import load_dotenv
    load_dotenv()
except ImportError:
    pass

try:
    from flask import Flask, Response, jsonify, request, send_file
except ImportError:
    Flask = None
    Response = None
    jsonify = None
    request = None
    send_file = None

# นำเข้า APScheduler สำหรับทำ Background Task ดึงข้อมูลจาก Google Sheets ตามรอบเวลา
try:
    from apscheduler.schedulers.background import BackgroundScheduler
    HAS_APSCHEDULER = True
except ImportError:
    BackgroundScheduler = None
    HAS_APSCHEDULER = False

# นำเข้าโมดูลประมวลผลข้อมูลและสร้างแผนที่/แดชบอร์ดจาก create_detailed_map.py
import create_detailed_map as cdm

# ตั้งค่า Encoding สำหรับ Server / Windows
if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
if hasattr(sys.stderr, "reconfigure"):
    sys.stderr.reconfigure(encoding="utf-8", errors="replace")

BASE_DIR = os.path.dirname(os.path.abspath(__file__)) if "__file__" in globals() else os.getcwd()
IS_RENDER_SERVER = bool(
    os.environ.get("RENDER")
    or os.environ.get("RENDER_SERVICE_ID")
    or (os.name != "nt" and os.path.isdir("/tmp"))
)

# ระยะเวลาในการอัปเดตข้อมูลเบื้องหลัง (ค่าเริ่มต้นทุกๆ 5 นาที = 300 วินาที ป้องกันการติด Rate Limit ของ Google Sheets)
REFRESH_INTERVAL_MINUTES = max(1, int(os.environ.get("SCADA_REFRESH_MINUTES", "5")))
REFRESH_INTERVAL_SECONDS = REFRESH_INTERVAL_MINUTES * 60

# ==============================================================================
# --- ส่วนจัดการ In-Memory Cache & Background Task (APScheduler) ---
# ==============================================================================
CACHE_LOCK = threading.Lock()
REFRESH_IN_PROGRESS_LOCK = threading.Lock()

# ตัวแปร Global In-Memory Cache สำหรับเก็บข้อมูลล่าสุดที่ดึงและประมวลผลจาก Google Sheets
GLOBAL_SCADA_CACHE = {
    "initialized": False,
    "is_updating": False,
    "version": 0,
    "last_updated": 0.0,
    "last_error": "",
    "overall_summary": {
        "total": 0,
        "Online": 0,
        "Connecting": 0,
        "Offline": 0,
        "Initializing": 0,
        "Telemetry Failure": 0,
        "other": 0,
        "offline_total": 0,
        "ava_pct": 0.0,
    },
    "dash_all_data": [],
    "search_data": [],
    "export_data": [],
    "report_data": [],
    "status_hash_map": {},
    "comm_device_groups": cdm.COMM_DEVICE_GROUPS,
    "display_fields": [],
    "html": "",
}

_SCHEDULER_INSTANCE = None
_SCHEDULER_START_LOCK = threading.Lock()


def refresh_scada_data_job(prefer_local=False):
    """
    [BACKGROUND TASK] ฟังก์ชันสำหรับดึงข้อมูลจาก Google Sheets, อัปเดตฐานข้อมูล SQLite,
    คำนวณค่าสรุปภาพรวม (overall_summary) และบันทึกลงใน In-Memory Cache (`GLOBAL_SCADA_CACHE`)
    """
    if not REFRESH_IN_PROGRESS_LOCK.acquire(blocking=False):
        print("[INFO] มีงานอัปเดตข้อมูลกำลังทำงานอยู่แล้ว ข้ามรอบซ้ำซ้อนนี้...")
        return False

    with CACHE_LOCK:
        GLOBAL_SCADA_CACHE["is_updating"] = True

    try:
        print(f"[SCHEDULER] เริ่มซิงค์ข้อมูลจาก Google Sheets และประมวลผล (prefer_local={prefer_local})...")
        write_static = (not IS_RENDER_SERVER) or (os.environ.get("SCADA_WRITE_STATIC_HTML") == "1")
        result = cdm.build_scada_map_and_data(prefer_local=prefer_local, write_files=write_static, embed_initial_data=True)

        with CACHE_LOCK:
            GLOBAL_SCADA_CACHE.update({
                "initialized": True,
                "is_updating": False,
                "version": result["version"],
                "last_updated": result["last_updated"],
                "last_error": "",
                "overall_summary": result["overall_summary"],
                "dash_all_data": result["dash_all_data"],
                "search_data": result["search_data"],
                "export_data": result["export_data"],
                "report_data": result["report_data"],
                "status_hash_map": result["status_hash_map"],
                "comm_device_groups": result["comm_device_groups"],
                "display_fields": result["display_fields"],
                "html": result["html"],
            })

        print(
            f"[SCHEDULER] อัปเดต In-Memory Cache สำเร็จ! (Version: {result['version']}, "
            f"จำนวนอุปกรณ์ทั้งระบบ: {result['overall_summary'].get('total', 0)} ชุด, "
            f"%AVA รวม: {result['overall_summary'].get('ava_pct', 0.0)}%)"
        )
        return True
    except Exception as exc:
        err_msg = str(exc)
        print(f"[SCHEDULER ERROR] เกิดข้อผิดพลาดในการอัปเดตข้อมูลเบื้องหลัง: {err_msg}")
        with CACHE_LOCK:
            GLOBAL_SCADA_CACHE["is_updating"] = False
            GLOBAL_SCADA_CACHE["last_error"] = err_msg
        return False
    finally:
        REFRESH_IN_PROGRESS_LOCK.release()


def get_cached_scada_snapshot():
    """อ่านข้อมูลล่าสุดจาก In-Memory Cache อย่างปลอดภัย (Thread-safe)"""
    with CACHE_LOCK:
        return dict(GLOBAL_SCADA_CACHE)


def build_api_map_data_payload():
    """
    [API PAYLOAD] สร้างโครงสร้าง JSON สำหรับ Endpoint `/api/map_data`
    ประกอบด้วย overall_summary (ยอดรวมทั้งระบบ), พิกัดแผนที่ Lat/Long (search_data / stations),
    รายละเอียดสถานะอุปกรณ์ (dash_all_data, export_data) และรายการตรวจสอบหน้างาน (report_data)
    """
    snap = get_cached_scada_snapshot()
    return {
        "success": snap["initialized"],
        "version": snap["version"],
        "last_updated": snap["last_updated"],
        "is_updating": snap["is_updating"],
        "error": snap["last_error"],
        "overall_summary": snap["overall_summary"],
        "stations": snap["search_data"],
        "search_data": snap["search_data"],
        "dash_all_data": snap["dash_all_data"],
        "export_data": snap["export_data"],
        "report_data": snap["report_data"],
        "status_hash_map": snap["status_hash_map"],
        "comm_device_groups": snap["comm_device_groups"],
        "display_fields": snap["display_fields"],
    }


def render_page_html_from_cache(view_mode="map", open_ai=False):
    """
    [CLIENT-SIDE RENDERING SHELL]
    เสิร์ฟโครงสร้าง HTML/JS Shell น้ำหนักเบา (หรือไฟล์ map.html) ให้กับ Browser ทันที
    โดยไม่ต้องประกอบ Marker แผนที่หรือตารางบนฝั่ง Backend (Browser จะเรียก `/api/map_data` เพื่อ Render เอง 100%)
    """
    snap = get_cached_scada_snapshot()
    html_content = snap.get("html") or ""
    if not html_content:
        static_map_path = os.path.join(BASE_DIR, "map.html")
        if os.path.exists(static_map_path):
            with open(static_map_path, "r", encoding="utf-8") as f:
                html_content = f.read()
        else:
            return None

    if view_mode == "dashboard":
        html_content = html_content.replace(
            "var dashDemoMode = false;",
            "window.DEFAULT_START_DASHBOARD = true;\nvar dashDemoMode = false;",
            1,
        )
    if open_ai:
        html_content = html_content.replace(
            "var dashDemoMode = false;",
            "window.DEFAULT_OPEN_AI_MODAL = true;\nvar dashDemoMode = false;",
            1,
        )
    return html_content


def start_background_scheduler():
    """
    [BACKGROUND TASK INITIALIZATION]
    1. โหลดข้อมูลตั้งต้นจากฐานข้อมูลในเครื่อง (SQLite / GIS Master) เข้าสู่ Memory ทันทีเพื่อให้ Server พร้อมตอบสนองทันทีที่เปิด
    2. สั่งซิงค์ข้อมูลล่าสุดจาก Google Sheets ใน Thread เบื้องหลังทันที 1 รอบ
    3. ตั้งเวลา `BackgroundScheduler` (APScheduler) ให้ดึงข้อมูลจาก Google Sheets อัตโนมัติทุกๆ 5 นาที
       (หากสภาพแวดล้อมใดไม่ได้ติดตั้งไลบรารี apscheduler จะสลับไปใช้ Daemon Thread Loop สำรองโดยอัตโนมัติ)
    """
    global _SCHEDULER_INSTANCE
    with _SCHEDULER_START_LOCK:
        if _SCHEDULER_INSTANCE is not None:
            return _SCHEDULER_INSTANCE

        # 1. โหลดข้อมูลตั้งต้นแบบรวดเร็ว (Fast Warm-up จาก Local SQLite / Master Data)
        snap = get_cached_scada_snapshot()
        if not snap["initialized"]:
            print("[INIT] กำลังเตรียมข้อมูลตั้งต้นเข้าสู่ In-Memory Cache (Fast Warm-up)...")
            refresh_scada_data_job(prefer_local=True)

        # 2. สั่งดึงข้อมูลสดจาก Google Sheets ใน Background Thread ทันทีโดยไม่บล็อกการบูตของ Render
        threading.Thread(
            target=refresh_scada_data_job,
            kwargs={"prefer_local": False},
            daemon=True,
            name="InitialGoogleSheetsSync",
        ).start()

        # 3. ตั้งเวลา BackgroundScheduler ทุกๆ REFRESH_INTERVAL_MINUTES นาที
        if HAS_APSCHEDULER and BackgroundScheduler is not None:
            scheduler = BackgroundScheduler(daemon=True, timezone="UTC")
            scheduler.add_job(
                func=refresh_scada_data_job,
                trigger="interval",
                minutes=REFRESH_INTERVAL_MINUTES,
                kwargs={"prefer_local": False},
                id="scada_google_sheets_sync_job",
                max_instances=1,
                coalesce=True,
                replace_existing=True,
            )
            scheduler.start()
            _SCHEDULER_INSTANCE = scheduler
            print(f"[SCHEDULER] เริ่มทำงาน APScheduler (BackgroundScheduler) อัปเดตข้อมูลทุกๆ {REFRESH_INTERVAL_MINUTES} นาที")
        else:
            def _fallback_periodic_loop():
                while True:
                    time.sleep(REFRESH_INTERVAL_SECONDS)
                    refresh_scada_data_job(prefer_local=False)

            t = threading.Thread(target=_fallback_periodic_loop, daemon=True, name="FallbackPeriodicSync")
            t.start()
            _SCHEDULER_INSTANCE = t
            print(
                f"[SCHEDULER] ไม่พบไลบรารี apscheduler -> ใช้ Background Thread ตั้งเวลาอัปเดตทุกๆ {REFRESH_INTERVAL_MINUTES} นาทีแทน"
            )

        return _SCHEDULER_INSTANCE


# ==============================================================================
# --- ส่วนจัดการ Web Server & API Endpoints (Flask + Fallback HTTP Server) ---
# ==============================================================================
class _FallbackFlask:
    """เซิร์ฟเวอร์สำรองมาตรฐาน (Built-in HTTP Server) กรณีรันเครื่อง Local ที่ยังไม่ได้ติดตั้งไลบรารี Flask"""
    def __init__(self, name):
        self.name = name
        self._routes = {}

    def route(self, path, **kwargs):
        def decorator(fn):
            self._routes[path] = fn
            return fn
        return decorator

    def run(self, host="0.0.0.0", port=10000, **kwargs):
        from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
        from urllib.parse import parse_qs, urlparse

        print(f"[WARNING] ไม่พบไลบรารี Flask กำลังรันด้วย Built-in HTTP Server สำรองที่ {host}:{port}...")

        class _Handler(BaseHTTPRequestHandler):
            def _send_json(self_req, data_obj, status_code=200):
                body = json.dumps(data_obj, ensure_ascii=False).encode("utf-8")
                self_req.send_response(status_code)
                self_req.send_header("Content-Type", "application/json; charset=utf-8")
                self_req.send_header("Cache-Control", "no-store, no-cache, must-revalidate, max-age=0")
                self_req.end_headers()
                self_req.wfile.write(body)

            def _send_static_file(self_req, filename, content_type):
                fpath = os.path.join(BASE_DIR, filename)
                if not os.path.exists(fpath):
                    self_req.send_response(404)
                    self_req.end_headers()
                    return
                with open(fpath, "rb") as f:
                    data = f.read()
                self_req.send_response(200)
                self_req.send_header("Content-Type", content_type)
                self_req.end_headers()
                self_req.wfile.write(data)

            def do_GET(self_req):
                parsed = urlparse(self_req.path)
                path = parsed.path
                qs = parse_qs(parsed.query)

                if path == "/healthz":
                    snap = get_cached_scada_snapshot()
                    self_req._send_json({
                        "status": "ok",
                        "initialized": snap["initialized"],
                        "version": snap["version"],
                        "last_updated": snap["last_updated"],
                        "time": time.time(),
                    })
                elif path == "/api/version":
                    snap = get_cached_scada_snapshot()
                    self_req._send_json({
                        "version": snap["version"],
                        "last_updated": snap["last_updated"],
                        "is_updating": snap["is_updating"],
                    })
                elif path in ("/api/map_data", "/api/map-data"):
                    if not get_cached_scada_snapshot()["initialized"]:
                        refresh_scada_data_job(prefer_local=True)
                    self_req._send_json(build_api_map_data_payload())
                elif path == "/api/gemini/status":
                    self_req._send_json(cdm.sync_and_get_session_state())
                elif path in ("/", "/index.html", "/portal"):
                    if os.path.exists(os.path.join(BASE_DIR, "index.html")):
                        self_req._send_static_file("index.html", "text/html; charset=utf-8")
                    else:
                        self_req.send_response(404)
                        self_req.end_headers()
                elif path in ("/procurement", "/procurement.html"):
                    self_req._send_static_file("procurement.html", "text/html; charset=utf-8")
                elif path in ("/logo.webp", "/logo.jpg"):
                    target = path.lstrip("/")
                    if os.path.exists(os.path.join(BASE_DIR, target)):
                        mime = "image/webp" if target.endswith(".webp") else "image/jpeg"
                        self_req._send_static_file(target, mime)
                    elif os.path.exists(os.path.join(BASE_DIR, "logo.webp")):
                        self_req._send_static_file("logo.webp", "image/webp")
                    else:
                        self_req._send_static_file("logo.jpg", "image/jpeg")
                elif path in (
                    "/map",
                    "/map.html",
                    "/scada_detailed_map.html",
                    "/dashboard",
                    "/dashboard.html",
                    "/scada_dashboard.html",
                    "/ai-report",
                    "/map-data",
                ):
                    if not get_cached_scada_snapshot()["initialized"]:
                        refresh_scada_data_job(prefer_local=True)
                    is_dash = path in ("/dashboard", "/dashboard.html", "/scada_dashboard.html") or qs.get("view", [""])[0] == "dashboard"
                    view_mode = "dashboard" if is_dash else "map"
                    open_ai = (path == "/ai-report") or (qs.get("ai", [""])[0] == "1") or (qs.get("open_ai", [""])[0] == "1")
                    html_txt = render_page_html_from_cache(view_mode=view_mode, open_ai=open_ai)
                    if not html_txt:
                        self_req.send_response(503)
                        self_req.end_headers()
                        self_req.wfile.write("กำลังจัดเตรียมข้อมูลระบบ SCADA...".encode("utf-8"))
                        return
                    body = html_txt.encode("utf-8")
                    self_req.send_response(200)
                    self_req.send_header("Content-Type", "text/html; charset=utf-8")
                    self_req.end_headers()
                    self_req.wfile.write(body)
                else:
                    self_req.send_response(404)
                    self_req.end_headers()

            def do_POST(self_req):
                parsed = urlparse(self_req.path)
                length = int(self_req.headers.get("Content-Length", 0))
                raw_body = self_req.rfile.read(length) if length > 0 else b"{}"
                try:
                    payload = json.loads(raw_body.decode("utf-8", errors="replace"))
                except Exception:
                    payload = {}

                if parsed.path in ("/api/user_report", "/api/user-report"):
                    res_ur = save_and_forward_user_report(payload)
                    self_req._send_json(res_ur, 200)
                elif parsed.path == "/api/gemini/summarize":
                    try:
                        res = cdm.summarize_inspection_report_with_gemini(
                            rows=payload.get("rows") or [],
                            preferred_model=payload.get("model") or "gemini-3.5-pro-exp",
                            api_key=str(payload.get("api_key") or "").strip() or cdm.DEFAULT_GEMINI_API_KEY,
                            custom_prompt=payload.get("custom_prompt"),
                        )
                        self_req._send_json(res, 200)
                    except Exception as exc:
                        self_req._send_json({
                            "success": False,
                            "error": str(exc),
                            "session_state": cdm.sync_and_get_session_state(),
                        }, 429)
                else:
                    self_req.send_response(404)
                    self_req.end_headers()

        srv = ThreadingHTTPServer((host, port), _Handler)
        srv.serve_forever()


app = Flask(__name__) if Flask is not None else _FallbackFlask(__name__)


@app.route("/api/map_data", methods=["GET"])
@app.route("/api/map-data", methods=["GET"])
def api_map_data():
    """
    [FEATURE 4] API Endpoint สำหรับ Seamless Map & Dashboard Update
    คืนค่าข้อมูลทั้งหมดจาก In-Memory Cache ในรูปแบบ JSON:
    - overall_summary: ค่าสรุปภาพรวมของทั้งระบบ
    - stations / search_data: ข้อมูลพิกัดแผนที่ (Lat/Long) และสถานะของแต่ละสถานี
    - dash_all_data / export_data: ข้อมูลรายละเอียดอุปกรณ์และการสื่อสารทั้งหมด
    - report_data: ข้อมูลประวัติการเข้าตรวจสอบและแก้ไขหน้างาน
    """
    if not get_cached_scada_snapshot()["initialized"]:
        refresh_scada_data_job(prefer_local=True)

    if request is not None and request.args.get("refresh") == "1":
        threading.Thread(
            target=refresh_scada_data_job,
            kwargs={"prefer_local": False},
            daemon=True,
        ).start()

    payload = build_api_map_data_payload()
    resp = jsonify(payload)
    resp.headers["Cache-Control"] = "no-store, no-cache, must-revalidate, max-age=0"
    return resp


@app.route("/")
@app.route("/index.html")
@app.route("/portal")
def portal_home():
    """
    Route หลักสำหรับเสิร์ฟหน้า Web Portal ของหน่วยงาน (index.html)
    หากเปิดด้วยพารามิเตอร์ ?view=dashboard หรือ ?ai=1 จะสลับไปเปิดหน้า SCADA ให้ทันที
    """
    if request is not None and (
        request.args.get("view") in ("dashboard", "map")
        or request.args.get("ai") == "1"
        or request.args.get("open_ai") == "1"
    ):
        return scada_app()

    portal_path = os.path.join(BASE_DIR, "index.html")
    if os.path.exists(portal_path):
        return send_file(portal_path, mimetype="text/html; charset=utf-8")

    # Fallback ไปยังหน้าแผนที่ SCADA กรณีไม่พบไฟล์ index.html
    return scada_app()


@app.route("/procurement")
@app.route("/procurement.html")
def procurement_page():
    """Route สำหรับเสิร์ฟหน้าระบบบริหารและติดตามงานจัดซื้อจัดจ้าง (procurement.html)"""
    proc_path = os.path.join(BASE_DIR, "procurement.html")
    if os.path.exists(proc_path):
        return send_file(proc_path, mimetype="text/html; charset=utf-8")
    return "ไม่พบไฟล์ procurement.html", 404


@app.route("/logo.webp")
@app.route("/logo.jpg")
def serve_logo():
    """Route สำหรับเสิร์ฟไฟล์โลโก้หน่วยงานบนหน้า Web Portal"""
    req_file = request.path.lstrip("/")
    file_path = os.path.join(BASE_DIR, req_file)
    if os.path.exists(file_path):
        mime = "image/webp" if req_file.endswith(".webp") else "image/jpeg"
        return send_file(file_path, mimetype=mime)

    for fallback_name, mime in [("logo.webp", "image/webp"), ("logo.jpg", "image/jpeg")]:
        fb_path = os.path.join(BASE_DIR, fallback_name)
        if os.path.exists(fb_path):
            return send_file(fb_path, mimetype=mime)

    return "ไม่พบไฟล์โลโก้", 404


@app.route("/map")
@app.route("/map.html")
@app.route("/scada_detailed_map.html")
@app.route("/dashboard")
@app.route("/dashboard.html")
@app.route("/scada_dashboard.html")
@app.route("/ai-report")
@app.route("/map-data")
def scada_app():
    """
    [FEATURE 4] Route สำหรับเสิร์ฟหน้าเว็บแผนที่และแดชบอร์ด SCADA (%AVA)
    รองรับการคลิกจากหน้า Web Portal (map.html / dashboard.html)
    โดยฝั่ง Frontend จะเรียกใช้ Fetch API (`/api/map_data`) เพื่อดึงข้อมูลมาอัปเดตแบบไม่รีเฟรชหน้า (Seamless Update)
    """
    if not get_cached_scada_snapshot()["initialized"]:
        refresh_scada_data_job(prefer_local=True)

    is_dash = (
        request.path in ("/dashboard", "/dashboard.html", "/scada_dashboard.html")
        or request.args.get("view") == "dashboard"
    )
    open_ai = (
        request.path.endswith("/ai-report")
        or request.args.get("ai") == "1"
        or request.args.get("open_ai") == "1"
    )

    html_content = render_page_html_from_cache(
        view_mode="dashboard" if is_dash else "map",
        open_ai=open_ai,
    )
    if not html_content:
        return (
            "<!DOCTYPE html><html lang='th'><head><meta charset='utf-8'>"
            "<meta http-equiv='refresh' content='3'>"
            "<title>กำลังเตรียมข้อมูลระบบ SCADA...</title></head>"
            "<body style='font-family:sans-serif;display:flex;justify-content:center;align-items:center;height:100vh;background:#f0f2f5;'>"
            "<div style='background:#fff;padding:24px 32px;border-radius:12px;box-shadow:0 4px 16px rgba(0,0,0,0.08);text-align:center;'>"
            "<h2 style='color:#1a73e8;margin:0 0 8px;'>กำลังเตรียมข้อมูลระบบ SCADA...</h2>"
            "<p style='color:#5f6368;margin:0;'>ระบบกำลังซิงค์ข้อมูลเริ่มต้น กรุณารอสักครู่</p>"
            "</div></body></html>",
            503,
        )

    if Response is not None:
        return Response(html_content, mimetype="text/html; charset=utf-8")
    return html_content


@app.route("/api/version", methods=["GET"])
def api_version():
    """คืนค่าเลขเวอร์ชันและเวลาอัปเดตล่าสุดของข้อมูลใน In-Memory Cache"""
    snap = get_cached_scada_snapshot()
    return jsonify({
        "version": snap["version"],
        "last_updated": snap["last_updated"],
        "is_updating": snap["is_updating"],
    })


@app.route("/healthz", methods=["GET"])
def healthz():
    """Health Check Endpoint สำหรับ Render"""
    snap = get_cached_scada_snapshot()
    return jsonify({
        "status": "ok",
        "initialized": snap["initialized"],
        "version": snap["version"],
        "last_updated": snap["last_updated"],
        "time": time.time(),
    })


def save_and_forward_user_report(payload):
    """บันทึกข้อมูลแจ้งแก้ไขอุปกรณ์ (User_Report) ลง SQLite และส่งต่อไปยัง Google Apps Script Webhook (หากตั้งค่าไว้)"""
    import sqlite3
    import urllib.request

    site_id = str(payload.get("site_id") or "-").strip()
    cmd = str(payload.get("cmd") or "-").strip()
    location = str(payload.get("location") or "-").strip()
    pea = str(payload.get("pea") or "-").strip()
    device_type = str(payload.get("device_type") or "-").strip()
    incorrect_fields = payload.get("incorrect_fields") or []
    corrected_value = str(payload.get("corrected_value") or "").strip()
    per_field_detail = str(payload.get("per_field_detail") or "").strip()
    reporter_note = str(payload.get("reporter_note") or "").strip()
    ts = str(payload.get("timestamp") or time.strftime("%Y-%m-%dT%H:%M:%SZ", time.gmtime()))

    saved_local = False
    try:
        import tempfile, shutil
        work_db = cdm.DB_PATH
        use_tmp = False
        try:
            tc = sqlite3.connect(work_db)
            tc.execute("CREATE TABLE IF NOT EXISTS _ping (id INT)")
            tc.close()
        except Exception:
            use_tmp = True
            work_db = os.path.join(tempfile.gettempdir(), "scada_real_work.db")
        conn = sqlite3.connect(work_db)
        conn.execute(
            """
            CREATE TABLE IF NOT EXISTS user_reports (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                timestamp TEXT,
                site_id TEXT,
                cmd TEXT,
                location TEXT,
                pea TEXT,
                device_type TEXT,
                incorrect_fields TEXT,
                corrected_value TEXT,
                per_field_detail TEXT,
                reporter_note TEXT,
                raw_json TEXT
            )
            """
        )
        conn.execute(
            """
            INSERT INTO user_reports (
                timestamp, site_id, cmd, location, pea, device_type,
                incorrect_fields, corrected_value, per_field_detail, reporter_note, raw_json
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                ts,
                site_id,
                cmd,
                location,
                pea,
                device_type,
                ", ".join(str(x) for x in incorrect_fields),
                corrected_value,
                per_field_detail,
                reporter_note,
                json.dumps(payload, ensure_ascii=False),
            ),
        )
        conn.commit()
        conn.close()
        if use_tmp and os.path.exists(work_db):
            try:
                shutil.copy2(work_db, cdm.DB_PATH)
            except Exception:
                pass
        saved_local = True
    except Exception as db_exc:
        print(f"[USER_REPORT DB WARNING] {db_exc}")

    forwarded = False
    webhook_url = str(payload.get("webhook_url") or cdm.USER_REPORT_GAS_WEBHOOK_URL or "").strip()
    if webhook_url:
        try:
            req_bytes = json.dumps(payload, ensure_ascii=False).encode("utf-8")
            req = urllib.request.Request(
                webhook_url,
                data=req_bytes,
                headers={"Content-Type": "text/plain;charset=utf-8"},
                method="POST",
            )
            opener = cdm._build_url_opener()
            with opener.open(req, timeout=15) as _:
                forwarded = True
        except Exception as wb_exc:
            print(f"[USER_REPORT WEBHOOK WARNING] {wb_exc}")

    return {
        "success": True,
        "saved_local": saved_local,
        "forwarded_to_webhook": forwarded,
        "site_id": site_id,
        "timestamp": ts,
    }


@app.route("/api/user_report", methods=["GET", "POST"])
@app.route("/api/user-report", methods=["GET", "POST"])
def api_user_report():
    """Endpoint สำหรับรับและบันทึกข้อมูลแจ้งแก้ไขอุปกรณ์ไม่ถูกต้อง (User_Report)"""
    if request.method == "GET":
        import sqlite3
        rows_out = []
        try:
            conn = sqlite3.connect(cdm.DB_PATH)
            cur = conn.cursor()
            cur.execute("SELECT name FROM sqlite_master WHERE type='table' AND name='user_reports'")
            if cur.fetchone():
                cur.execute(
                    "SELECT id, timestamp, site_id, cmd, location, pea, device_type, incorrect_fields, corrected_value, per_field_detail, reporter_note FROM user_reports ORDER BY id DESC LIMIT 100"
                )
                cols = [d[0] for d in cur.description]
                rows_out = [dict(zip(cols, r)) for r in cur.fetchall()]
            conn.close()
        except Exception:
            pass
        return jsonify({"success": True, "reports": rows_out})

    payload = request.get_json(silent=True) or {}
    if not payload.get("site_id"):
        return jsonify({"success": False, "error": "กรุณาระบุ Site ID ของอุปกรณ์"}), 400
    res = save_and_forward_user_report(payload)
    return jsonify(res)


@app.route("/api/gemini/status", methods=["GET"])
def api_gemini_status():
    """คืนค่าสถานะโควต้าและ Cooldown ของโมเดล Gemini AI ทั้งหมด"""
    return jsonify(cdm.sync_and_get_session_state())


@app.route("/api/gemini/simulate-limit", methods=["POST"])
def api_gemini_simulate_limit():
    """Endpoint สำหรับจำลองหรือรีเซ็ตสถานะ Rate Limit (429) / RPD ของโมเดล Gemini AI"""
    payload = request.get_json(silent=True) or {}
    action = payload.get("action", "reset")
    target_model = payload.get("model", "gemini-3.5-pro-exp")
    if action == "reset":
        cdm.reset_model_limit_state(target_model if target_model != "all" else None)
    elif action == "rate_limit_429":
        cdm.mark_model_rate_limited(
            target_model,
            error_msg="Simulated 429 Too Many Requests (Rolling Window 60s)",
            is_rpd=False,
        )
    elif action == "rpd_full":
        cdm.mark_model_rate_limited(
            target_model,
            error_msg="Simulated RPD Daily Quota Exhausted (Reset 07:00 TH)",
            is_rpd=True,
        )
    return jsonify(cdm.sync_and_get_session_state())


@app.route("/api/gemini/summarize", methods=["POST"])
def api_gemini_summarize():
    """Endpoint สำหรับสรุปและเรียบเรียงรายงานการตรวจสอบอุปกรณ์ด้วย Google Gemini API"""
    payload = request.get_json(silent=True) or {}
    rows = payload.get("rows") or []
    preferred_model = payload.get("model") or "gemini-3.5-pro-exp"
    custom_api_key = str(payload.get("api_key") or "").strip() or cdm.DEFAULT_GEMINI_API_KEY

    if not rows:
        return jsonify({
            "success": False,
            "error": "ไม่พบข้อมูลสำหรับการสรุปรายงาน",
            "session_state": cdm.sync_and_get_session_state(),
        }), 400

    try:
        result = cdm.summarize_inspection_report_with_gemini(
            rows=rows,
            preferred_model=preferred_model,
            api_key=custom_api_key,
            custom_prompt=payload.get("custom_prompt"),
        )
        return jsonify(result)
    except Exception as exc:
        return jsonify({
            "success": False,
            "error": str(exc),
            "session_state": cdm.sync_and_get_session_state(),
        }), 429


@app.route("/api/export", methods=["POST"])
@app.route("/api/export-report", methods=["POST"])
def api_export_report():
    """Endpoint สำหรับสร้างและส่งออกไฟล์รายงาน Excel (.xlsx) ขนาดมาตรฐาน A4 แนวนอน"""
    payload = request.get_json(silent=True) or {}
    rows = payload.get("rows") or []
    preferred_model = payload.get("model") or "gemini-3.5-pro-exp"
    use_ai = bool(payload.get("use_ai", True))
    include_worker_col = bool(payload.get("include_worker_col", False))
    response_format = payload.get("format", "xlsx")
    custom_api_key = str(payload.get("api_key") or "").strip() or cdm.DEFAULT_GEMINI_API_KEY

    if not rows:
        return jsonify({
            "success": False,
            "error": "ไม่พบรายการข้อมูลสำหรับส่งออกรายงาน",
            "session_state": cdm.sync_and_get_session_state(),
        }), 400

    if use_ai:
        try:
            ai_result = cdm.summarize_inspection_report_with_gemini(
                rows=rows,
                preferred_model=preferred_model,
                api_key=custom_api_key,
                custom_prompt=payload.get("custom_prompt"),
            )
        except Exception as exc:
            base_exec = cdm._build_deterministic_executive_summary(rows)
            rewritten = [dict(r) for r in rows]
            for r in rewritten:
                r["ai_detail"] = cdm._deterministic_categorize_row_detail(
                    r.get("detail", ""), r.get("resp", ""), r.get("status", "")
                )
                r["detail"] = r["ai_detail"]
            ai_result = {
                "success": False,
                "warning": str(exc),
                "used_model": preferred_model,
                "chunk_count": 1,
                "rows": rewritten,
                "summary_markdown": base_exec["summary_markdown"],
                "categories": base_exec["categories"],
                "metrics": base_exec["metrics"],
                "session_state": cdm.sync_and_get_session_state(),
            }
    else:
        base_exec = cdm._build_deterministic_executive_summary(rows)
        rewritten = [dict(r) for r in rows]
        for r in rewritten:
            r["ai_detail"] = r.get("ai_detail") or cdm._deterministic_categorize_row_detail(
                r.get("detail", ""), r.get("resp", ""), r.get("status", "")
            )
        ai_result = {
            "success": True,
            "warning": "",
            "used_model": preferred_model,
            "chunk_count": 1,
            "rows": rewritten,
            "summary_markdown": base_exec["summary_markdown"],
            "categories": base_exec["categories"],
            "metrics": base_exec["metrics"],
            "session_state": cdm.sync_and_get_session_state(),
        }

    if response_format == "json":
        return jsonify(ai_result)

    try:
        import openpyxl
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from openpyxl.utils import get_column_letter

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Report"
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.page_setup.orientation = ws.ORIENTATION_LANDSCAPE
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0

        thin_side = Side(style="thin", color="000000")
        cell_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)
        header_fill = PatternFill(start_color="F1F3F4", end_color="F1F3F4", fill_type="solid")

        headers = (
            [
                "ลำดับ",
                "รหัสสั่งการ\n(ข้อมูล GIS)",
                "Site ID",
                "สถานที่\n(ข้อมูล GIS)",
                "รายละเอียดการตรวจสอบ/แก้ไข\n(จัดหมวดหมู่โดย AI)",
                "ผู้รับผิดชอบและดำเนินการแก้ไขต่อไป",
                "ผู้ปฏิบัติงาน (ผอส.ฉ.2)",
                "สถานะหลังดำเนินการ",
            ]
            if include_worker_col
            else [
                "ลำดับ",
                "รหัสสั่งการ\n(ข้อมูล GIS)",
                "Site ID",
                "สถานที่\n(ข้อมูล GIS)",
                "รายละเอียดการตรวจสอบ/แก้ไข\n(จัดหมวดหมู่โดย AI)",
                "ผู้รับผิดชอบและดำเนินการแก้ไขต่อไป",
                "สถานะหลังดำเนินการ",
            ]
        )
        last_col_letter = get_column_letter(len(headers))
        ws.merge_cells(f"A1:{last_col_letter}1")
        ws["A1"] = "รายละเอียดการตรวจสอบ/แก้ไขอุปกรณ์ FDCU OFFLine (เรียบเรียงด้วย Gemini AI)"
        ws["A1"].font = Font(name="TH SarabunPSK", size=20, bold=True)
        ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 30

        ws.append(headers)
        ws.row_dimensions[2].height = 36
        for col_idx in range(1, len(headers) + 1):
            c = ws.cell(row=2, column=col_idx)
            c.font = Font(name="TH SarabunPSK", size=16, bold=True)
            c.fill = header_fill
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            c.border = cell_border

        final_rows = ai_result.get("rows") or rows
        for idx, r in enumerate(final_rows, start=1):
            detail_val = r.get("ai_detail") or r.get("detail") or "-"
            row_vals = (
                [
                    idx,
                    r.get("cmd", "-"),
                    r.get("site", "-"),
                    r.get("loc", "-"),
                    detail_val,
                    r.get("resp", "-"),
                    r.get("workers", "-"),
                    r.get("status", "-"),
                ]
                if include_worker_col
                else [
                    idx,
                    r.get("cmd", "-"),
                    r.get("site", "-"),
                    r.get("loc", "-"),
                    detail_val,
                    r.get("resp", "-"),
                    r.get("status", "-"),
                ]
            )
            ws.append(row_vals)
            cur_r = idx + 2
            for col_idx in range(1, len(row_vals) + 1):
                c = ws.cell(row=cur_r, column=col_idx)
                c.font = Font(name="TH SarabunPSK", size=16)
                c.border = cell_border
                if col_idx == 5 or (include_worker_col and col_idx == 7):
                    c.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
                else:
                    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        col_widths = [7, 18, 18, 25, 56, 28, 30, 20] if include_worker_col else [8, 20, 20, 28, 65, 32, 22]
        for i, w in enumerate(col_widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w

        bio = io.BytesIO()
        wb.save(bio)
        bio.seek(0)
        return send_file(
            bio,
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            as_attachment=True,
            download_name="SCADA_Inspection_Report_AI_A4.xlsx",
        )
    except Exception as exc:
        return jsonify({"success": False, "error": str(exc)}), 500


# สั่งเริ่มทำงาน Background Scheduler ทันทีเมื่อโมดูลถูกโหลด (รองรับทั้งการรันผ่าน Gunicorn บน Render และการรันผ่าน python app.py)
if "--static" not in sys.argv and "--generate-only" not in sys.argv:
    start_background_scheduler()


if __name__ == "__main__":
    import webbrowser

    if "--static" in sys.argv or "--generate-only" in sys.argv:
        cdm.build_scada_map_and_data(
            prefer_local=("--offline" in sys.argv or "--use-cache" in sys.argv),
            write_files=True,
        )
    else:
        # [FEATURE 2] รับค่า PORT จาก Environment Variable สำหรับการ Deploy บน Render และรันด้วย host="0.0.0.0"
        port = int(os.environ.get("PORT", 10000))
        if (
            not IS_RENDER_SERVER
            and "--no-browser" not in sys.argv
            and os.environ.get("WERKZEUG_RUN_MAIN") != "true"
        ):
            local_url = f"http://127.0.0.1:{port}"
            print(f"กำลังเปิดเว็บเบราว์เซอร์อัตโนมัติไปที่: {local_url} ...")
            threading.Timer(1.5, lambda: webbrowser.open(local_url)).start()
        app.run(host="0.0.0.0", port=port)